# -*- coding: utf-8 -*-
"""盘点模块

标准账实核对流程（对齐通用盘点单模板）：
    建盘点单（按账面数生成空白盘点表）
      → 现场盲盘填「实盘数」
      → 自动算「差异 = 实盘 − 账面」「差异率」
      → 填差异原因、处理意见
      → 确认「调整」，按差异生成盘盈入库 / 盘亏出库单据

关键设计：**盘点单本身不动库存**，只有最后「调整」那一步才生成单据动账。
好处是盘点期间可以继续收货发货，且差异原因、处理意见都留得下痕迹，
事后能查到"这批货为什么少了、谁批的、怎么处理的"。
"""

from flask import request, redirect, url_for, Response, render_template
from datetime import datetime
import io, csv, os, sqlite3
from urllib.parse import quote

from ..core.router import Router
bp = Router('count')

from ..core import db, util
from ..core.util import *            # noqa: F401,F403
from ..core.util import (_log_err, _last_prices, num, int_arg, safe_date,
                         today, clean_kw, ints)  # noqa: F401

# 差异原因与处理意见的标准选项（可手输其它值）
REASONS = ['', '出库漏录', '入库漏录', '多发少发', '串码/错发', '自然损耗',
           '破损/报废', '被盗丢失', '供应商多送未入账', '计量误差', '系统数据错误', '其它']
HANDLES = ['', '补录单据', '调整账面', '报损处理', '追责赔偿', '退回供应商', '待查', '其它']


# ---------------- 工具 ----------------
def _next_sno(d):
    """盘点单号：PD20260916-001（当天流水）

    用「当天最大序号 +1」而不是 COUNT+1：删过单子后 COUNT 会变小，
    重号会让 INSERT 撞上 UNIQUE 直接 500。
    另外两个请求可能同时算到同一个号，所以调用处还有重试兜底。
    """
    d = safe_date(d) or today()
    like = 'PD%s-%%' % d.replace('-', '')
    mx = db.q("SELECT MAX(sno) m FROM stk WHERE sno LIKE ?", like)[0]['m']
    n = 0
    if mx:
        try:
            n = int(str(mx)[-3:])
        except ValueError:
            n = 0
    return 'PD%s-%03d' % (d.replace('-', ''), n + 1)


def _items(stk_id):
    """盘点表明细（带物料信息、差异、差异率）"""
    rows = db.q("SELECT si.*, m.name, m.code, m.spec, m.width, m.unit, m.category,"
                " m.supplier"
                " FROM stk_items si JOIN materials m ON m.id=si.material_id"
                " WHERE si.stk_id=? ORDER BY si.sort, si.id", stk_id)
    out = []
    for r in rows:
        d = dict(r)
        b = float(d['book_qty'] or 0)
        rq = d['real_qty']
        d['counted'] = rq is not None
        if rq is None:
            d['diff'] = None
            d['rate'] = None
        else:
            d['diff'] = round(float(rq) - b, 6)
            d['rate'] = round(d['diff'] / b * 100, 2) if b else None
        d['diff_amt'] = round(d['diff'] * float(d['price'] or 0), 2) if d['diff'] else 0.0
        out.append(d)
    return out


def _summary(items):
    """汇总：物料种类、盘盈盘亏种类与数量、差异金额

    book   = 全部物料的账面合计（应盘）
    book_c = 只统计「已盘」物料的账面合计

    差异与差异率必须用 book_c 而不是 book：盘点进行到一半时，
    「实盘」只累加已盘行，若拿它去减「全部账面」，会把未盘物料的账面数
    当成盘亏显示出来。实测 250 种只盘了 100 种、且这 100 种全部盘盈 +3，
    合计却显示 -3946 的虚假大额亏损。
    """
    s = dict(n=len(items), done=0, over=0, short=0, match=0,
             over_qty=0.0, short_qty=0.0, amt=0.0, book=0.0, real=0.0,
             book_c=0.0, partial=False)
    for it in items:
        s['book'] += float(it['book_qty'] or 0)
        if not it['counted']:
            continue
        s['done'] += 1
        s['real'] += float(it['real_qty'] or 0)
        s['book_c'] += float(it['book_qty'] or 0)
        df = it['diff'] or 0
        if abs(df) < 1e-9:
            s['match'] += 1
        elif df > 0:
            s['over'] += 1
            s['over_qty'] += df
        else:
            s['short'] += 1
            s['short_qty'] += -df
        s['amt'] += it['diff_amt'] or 0
    # 差异、差异率都只按已盘部分算；没盘完时 partial=True，页面上要写明进度，
    # 免得"账面总数（全部）"和"实盘总数（已盘）"并排放着被误读成亏损。
    s['partial'] = s['done'] < s['n']
    s['rate'] = round((s['real'] - s['book_c']) / s['book_c'] * 100, 2) if s['book_c'] else None
    return s


# ---------------- 列表 ----------------
@bp.route('/stk')
def stk_list():
    kw = clean_kw(request.args.get('kw'))
    st = request.args.get('st') or ''
    w, a = [], []
    if kw:
        w.append("(sno LIKE ? OR scope LIKE ? OR counter LIKE ?)")
        a += ['%%%s%%' % kw] * 3
    if st:
        w.append("status=?")
        a.append(st)
    sql = "SELECT * FROM stk"
    if w:
        sql += " WHERE " + " AND ".join(w)
    sql += " ORDER BY sdate DESC, id DESC LIMIT 200"
    rows = db.q(sql, *a)
    for r in rows:
        r = dict(r)
    # 每张单子的差异概况
    info = {}
    for r in rows:
        it = _items(r['id'])
        s = _summary(it)
        info[r['id']] = s
    return render_template('stk_list.html', rows=rows, info=info, kw=kw, st=st,
                           statuses=['草稿', '已盘点', '已调整', '已取消'])


# ---------------- 新建 ----------------
@bp.route('/stk/new', methods=['GET', 'POST'])
def stk_new():
    if request.method == 'GET':
        tpls = db.q("SELECT id, name FROM tpl ORDER BY id")
        return render_template('stk_new.html', tpls=tpls, today=today(),
                               counters=_people('counter'), checkers=_people('checker'))
    d = safe_date(request.form.get('sdate')) or today()
    tpl_id = int_arg(request.form, 'tpl_id', 0)
    scope = (request.form.get('scope') or '').strip()[:60]
    counter = (request.form.get('counter') or '').strip()[:40]
    checker = (request.form.get('checker') or '').strip()[:40]
    note = (request.form.get('note') or '').strip()[:200]
    kw = clean_kw(request.form.get('kw'))
    with_zero = request.form.get('with_zero') == '1'

    sql = "SELECT * FROM v_stock WHERE 1=1"
    a = []
    if tpl_id:
        sql += " AND COALESCE(tpl_id,0)=?"
        a.append(tpl_id)
    if kw:
        sql += " AND (name LIKE ? OR code LIKE ? OR supplier LIKE ? OR category LIKE ?)"
        a += ['%%%s%%' % kw] * 4
    if not with_zero:
        sql += " AND stock<>0"
    sql += " ORDER BY name"
    mats = db.q(sql, *a)
    if not mats:
        return render_template('error.html', code=400, title='没有可盘点的物料',
                               detail='当前筛选条件下库存为空。可以勾选「包含零库存物料」再试。'), 400

    lp = _last_prices()
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    # 单号是「算出最大号再插」，两个人同时建单会撞号（UNIQUE）。
    # 撞了就换个号重试，而不是把 500 抛给用户。
    for _try in range(30):
        sno = _next_sno(d)
        try:
            with db.tx() as c:
                cur = c.execute("INSERT INTO stk(sno,sdate,tpl_id,scope,counter,checker,"
                                "status,note,created_at) VALUES(?,?,?,?,?,?,'草稿',?,?)",
                                (sno, d, tpl_id or None, scope, counter, checker, note, now))
                sid = cur.lastrowid
                for i, m in enumerate(mats):
                    c.execute("INSERT INTO stk_items(stk_id,material_id,book_qty,price,sort)"
                              " VALUES(?,?,?,?,?)",
                              (sid, m['id'], round(float(m['stock'] or 0), 6),
                               float(lp.get(m['id']) or 0), i))
            return redirect(url_for('stk_detail', sid=sid))
        except sqlite3.IntegrityError:
            continue
    return render_template('error.html', code=500, title='单号生成失败',
                           detail='请重试一次。'), 500


def _people(field):
    """以前填过的盘点人/监盘人，作为输入候选"""
    try:
        return [r[field] for r in db.q(
            "SELECT DISTINCT %s AS %s FROM stk WHERE %s<>'' ORDER BY %s LIMIT 12"
            % (field, field, field, field)) if r[field]]
    except Exception:
        return []


# ---------------- 明细 ----------------
@bp.route('/stk/<int:sid>')
def stk_detail(sid):
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单',
                               detail='它可能已被删除。'), 404
    row = row[0]
    items = _items(sid)
    return render_template('stk_detail.html', row=row, items=items,
                           summary=_summary(items), reasons=REASONS, handles=HANDLES,
                           unit_choices=db.unit_choices())


@bp.route('/stk/<int:sid>/save', methods=['POST'])
def stk_save(sid):
    """保存实盘数 / 原因 / 处理意见（不动库存）"""
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    row = row[0]
    if row['status'] in ('已调整', '已取消'):
        return redirect(url_for('stk_detail', sid=sid))
    ids = ints(request.form, 'item_id')
    n = 0
    with db.tx() as c:
        for iid in ids:
            rq = (request.form.get('real_%d' % iid) or '').strip()
            price = (request.form.get('price_%d' % iid) or '').strip()
            reason = (request.form.get('reason_%d' % iid) or '').strip()[:80]
            handle = (request.form.get('handle_%d' % iid) or '').strip()[:80]
            note = (request.form.get('note_%d' % iid) or '').strip()[:100]
            val = None
            if rq != '':
                val = num(rq, default=None, lo=0, hi=QTY_MAX)
                if val is None:
                    continue            # 填了但不是合法数字 -> 跳过这格
            c.execute("UPDATE stk_items SET real_qty=?, price=COALESCE(?,price),"
                      " reason=?, handle=?, note=? WHERE id=? AND stk_id=?",
                      (val, num(price, default=None) if price else None,
                       reason, handle, note, iid, sid))
            n += 1
        # 只要盘过至少一行就从「草稿」进到「已盘点」
        if row['status'] == '草稿':
            any_counted = db.q("SELECT COUNT(*) c FROM stk_items WHERE stk_id=?"
                               " AND real_qty IS NOT NULL", sid)[0]['c']
            if any_counted:
                c.execute("UPDATE stk SET status='已盘点' WHERE id=?", (sid,))
    return redirect(url_for('stk_detail', sid=sid) + '?saved=%d' % n)


# ---------------- 调整（生成单据，动库存） ----------------
@bp.route('/stk/<int:sid>/adjust', methods=['POST'])
def stk_adjust(sid):
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    row = row[0]
    if row['status'] == '已调整':
        return redirect(url_for('stk_detail', sid=sid))
    if not take_nonce(request.form.get('_nonce')):
        return redirect(url_for('stk_detail', sid=sid) + '?err=dup')

    items = [it for it in _items(sid) if it['counted'] and abs(it['diff'] or 0) > 1e-9]
    if not items:
        return redirect(url_for('stk_detail', sid=sid) + '?err=nodiff')

    d = row['sdate']
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    made = 0
    skipped = []
    with db.tx() as c:
        for it in items:
            diff = it['diff']
            kind = '进' if diff > 0 else '出'
            qty = abs(diff)
            # 盘亏不能超过现有库存（否则出库会把库存打成负数）
            if kind == '出':
                st = c.execute("SELECT COALESCE(SUM(CASE WHEN kind='进' THEN qty ELSE -qty END),0) s"
                               " FROM txns WHERE material_id=?", (it['material_id'],)).fetchone()
                avail = float(st['s'] or 0) + float(db.q("SELECT opening FROM materials WHERE id=?",
                                                         it['material_id'])[0]['opening'] or 0)
                if qty > avail + 1e-9:
                    skipped.append('%s（可用 %.2f）' % (it['name'], avail))
                    continue
            note = '%s调整 %s' % ('盘盈' if diff > 0 else '盘亏', row['sno'])
            if it['reason']:
                note += ' · ' + it['reason']
            cur = c.execute("INSERT INTO txns(tdate,material_id,kind,qty,price,note,"
                            "created_at,batch) VALUES(?,?,?,?,?,?,?,?)",
                            (d, it['material_id'], kind, round(qty, 6),
                             float(it['price'] or 0) or None, note[:200], now, ''))
            c.execute("UPDATE stk_items SET txn_id=? WHERE id=?", (cur.lastrowid, it['id']))
            made += 1
        # 只有真生成了单据才改状态：全部被拦截时不能假装「已调整」
        if made:
            c.execute("UPDATE stk SET status='已调整' WHERE id=?", (sid,))
    if skipped:
        url = url_for('stk_detail', sid=sid) + '?skip=' + quote('、'.join(skipped))
        if not made:
            url += '&none=1'
        return redirect(url)
    return redirect(url_for('stk_detail', sid=sid) + '?made=%d' % made)


@bp.route('/stk/<int:sid>/unadjust', methods=['POST'])
def stk_unadjust(sid):
    """撤销调整：删掉这次生成的单据，回到「已盘点」"""
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    if row[0]['status'] != '已调整':
        return redirect(url_for('stk_detail', sid=sid))
    with db.tx() as c:
        for r in db.q("SELECT id, txn_id FROM stk_items WHERE stk_id=? AND txn_id IS NOT NULL", sid):
            c.execute("DELETE FROM txns WHERE id=?", (r['txn_id'],))
            c.execute("UPDATE stk_items SET txn_id=NULL WHERE id=?", (r['id'],))
        c.execute("UPDATE stk SET status='已盘点' WHERE id=?", (sid,))
    return redirect(url_for('stk_detail', sid=sid) + '?undone=1')


@bp.route('/stk/<int:sid>/status', methods=['POST'])
def stk_status(sid):
    st = (request.form.get('status') or '').strip()
    if st in ('草稿', '已盘点', '已取消'):
        db.run("UPDATE stk SET status=? WHERE id=?", st, sid)
    return redirect(url_for('stk_detail', sid=sid))


@bp.route('/stk/<int:sid>/del', methods=['POST', 'GET'])
def stk_del(sid):
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    if row[0]['status'] == '已调整':
        # 调整过的单子不能直接删——先撤销调整，否则库存里会留下无主的调整单据
        return render_template('error.html', code=400, title='请先撤销调整',
                               detail='这张盘点单已经生成了盘盈/盘亏单据，'
                                      '直接删除会让库存对不上。请先点「撤销调整」再删除。'), 400
    with db.tx() as c:
        c.execute("DELETE FROM stk_items WHERE stk_id=?", (sid,))
        c.execute("DELETE FROM stk WHERE id=?", (sid,))
    return redirect(url_for('stk_list'))


# ---------------- 导出 ----------------
@bp.route('/stk/<int:sid>/export')
def stk_export(sid):
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    row = row[0]
    items = _items(sid)
    s = _summary(items)
    head = ['盘点单号', '盘点日期', '盘点范围', '物料编码', '物料名称', '规格', '单位',
            '账面数量', '实盘数量', '差异数量', '差异率(%)', '单价', '差异金额',
            '差异原因', '处理意见', '盘点人', '监盘人', '备注']
    rows = []
    for it in items:
        rows.append([row['sno'], row['sdate'], row['scope'] or '', it['code'] or '',
                     it['name'], it['spec'] or '', it['unit'] or '',
                     it['book_qty'], '' if it['real_qty'] is None else it['real_qty'],
                     '' if it['diff'] is None else it['diff'],
                     '' if it['rate'] is None else it['rate'],
                     it['price'] or '', it['diff_amt'] or '',
                     it['reason'] or '', it['handle'] or '', row['counter'] or '',
                     row['checker'] or '', it['note'] or ''])
    # 差异数量按「已盘部分」算：实盘只累加已盘行，若减全部账面，
    # 未盘物料的账面数会被当成盘亏显示出来（实测 250 种只盘 100 种时显示 -3946 的假亏损）。
    # 没盘完时在备注里写明进度，免得"账面（全部）"与"实盘（已盘）"并排被误读。
    tail = ('已盘 %d/%d 种 · 总差异率' % (s['done'], s['n'])) if s['partial'] else '总差异率'
    rows.append(['合计', '', '', '%d 种' % s['n'], '盘盈 %d 种 / 盘亏 %d 种' % (s['over'], s['short']),
                 '', '', s['book'], s['real'], round(s['real'] - s['book_c'], 2),
                 '' if s['rate'] is None else s['rate'], '', round(s['amt'], 2),
                 '', '', '', '', tail])

    fmt = (request.args.get('fmt') or 'xlsx').lower()
    if fmt == 'csv':
        return _csv(head, rows, row['sno'])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return _csv(head, rows, row['sno'])
    wb = Workbook()
    ws = wb.active
    ws.title = '盘点表'
    ws.append(head)
    hf = Font(bold=True, color='FFFFFF')
    hfill = PatternFill('solid', fgColor='4F6B8A')
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append(r)
    last = ws.max_row
    for c in ws[last]:
        c.font = Font(bold=True)
    # 差异不为 0 的标红
    red = Font(color='C0392B')
    for i in range(2, last):
        v = ws.cell(row=i, column=10).value
        try:
            if v not in (None, '') and abs(float(v)) > 1e-9:
                ws.cell(row=i, column=10).font = red
                ws.cell(row=i, column=11).font = red
        except (TypeError, ValueError):
            pass
    ws.freeze_panes = 'A2'
    for col, w in zip('ABCDEFGHIJKLMNOPQR',
                      [15, 11, 12, 12, 18, 12, 7, 10, 10, 10, 10, 9, 10, 16, 12, 9, 9, 14]):
        ws.column_dimensions[col].width = w
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fn = '盘点表_%s.xlsx' % row['sno']
    return Response(bio.getvalue(),
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':
                             "attachment; filename*=UTF-8''%s" % quote(fn)})


def _csv(head, rows, sno):
    bio = io.StringIO()
    w = csv.writer(bio)
    w.writerow(head)
    for r in rows:
        w.writerow(r)
    fn = '盘点表_%s.csv' % sno
    return Response('\ufeff' + bio.getvalue(), mimetype='text/csv; charset=utf-8',
                    headers={'Content-Disposition':
                             "attachment; filename*=UTF-8''%s" % quote(fn)})


# ---------------- 打印（现场盲盘用，只给物料和空实盘栏） ----------------
@bp.route('/stk/<int:sid>/print')
def stk_print(sid):
    row = db.q("SELECT * FROM stk WHERE id=?", sid)
    if not row:
        return render_template('error.html', code=404, title='找不到这张盘点单'), 404
    row = row[0]
    items = _items(sid)
    return render_template('stk_print.html', row=row, items=items, summary=_summary(items))
