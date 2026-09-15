"""从 Excel / CSV 批量导入物料档案。

表头自动识别（中文别名），按「料号 → 名称」优先级匹配已有物料：
  命中则更新非空白字段，未命中则新增。
"""
import io, csv, datetime, re

try:
    import openpyxl
except ImportError:
    openpyxl = None

# 规范字段 -> 可能出现的中文表头别名
ALIASES = {
    'name':     ['物料名称', '名称', '品名', '品名规格', '物料', '材料名称', 'name'],
    'code':     ['料号', '物料编号', '物料编码', '编号', '编码', '型号', '规格型号', 'code'],
    'supplier': ['供应商', '厂商', '供货商', 'supplier'],
    'category': ['类型', '类别', '分类', '大类', 'category'],
    'spec':     ['规格', '规格（米）', '规格(米)', '厚度', 'spec'],
    'width':    ['宽幅', '宽度', '幅宽', 'width'],
    'unit':     ['单位', '单位（卷）', '单位(卷)', '计量单位', 'unit'],
    'opening':  ['期初结存', '期初', '上月结存', '上期结存', '库存', '当前库存', 'opening'],
    'safety':   ['安全库存', '预警值', '库存预警', '最低库存', 'safety'],
    'status':   ['状态', '使用状态', 'status'],
}
NUM = {'opening', 'safety'}

# 单笔数量上限：挡住表里多输几个零（1e15）把库存冲爆的情况。
# 真有超大批量请拆单，或换个更大的单位（比如用吨而不是克）。
QTY_MAX = 1e9

# ---------- 流水（出入库单据）导入的表头别名 ----------
# 说明：物料的「类型」与「进出」易混淆，进出用 进出/方向/出入库 等专属词；
#      也支持 入库数量 / 出库数量 分列写法（一行可生成两笔）
TXN_ALIASES = {
    'date':  ['日期', '单据日期', '出入库日期', '入库日期', '出库日期', '时间', '业务日期', 'date'],
    'name':  ['物料名称', '名称', '品名', '物料', '材料名称', '货品名', 'name'],
    'code':  ['料号', '物料编号', '物料编码', '编号', '编码', '型号', '规格型号', 'code'],
    'kind':  ['进出', '进/出', '出入库', '方向', '收发明细', '出入库类型', '类型进销',
              '进出标志', 'in/out', 'type'],
    'qty':   ['数量', '出入库数量', '数量(平米)', '数量（平米）', '重量', '平米数', '收发数量', 'qty'],
    'in_qty':  ['进', '入库', '入库数量', '进货数量', '收入数量', '入库数', '进仓数量', 'in'],
    'out_qty': ['出', '出库', '出库数量', '出货数量', '发出数量', '出库数', '出仓数量', '领用数量', 'out'],
    'note':  ['备注', '客户', '单号', '用途', '领用人', '说明', '摘要', 'note'],
    'pieces':    ['件数', '件', '卷数', '箱数', '数量(件)', '数量（件）', 'pieces'],
    'per_piece': ['每件', '每件/个', '每件数量', '单件', '每件米数', '米/件', 'per'],
    'price': ['单价', '价格', '单位价格', '售价', '进货价', '单价(元)', '单价（元）',
              '元/米', '元/卷', '元/平', 'price'],
    'amount': ['金额', '总价', '合计金额', '总金额', '金额(元)', '金额（元）', 'amount'],
}
# 物料档案列（导入流水时顺带建档/更新用）
TXN_MAT_COLS = ['supplier', 'category', 'spec', 'width', 'unit', 'status', 'opening', 'safety']

def _norm(s):
    return str(s or '').strip().lower().replace(' ', '').replace('（', '(').replace('）', ')')

def unit_from_label(label, sample_vals):
    """「单位（卷）」这类表头，列里填的其实是卷数（50、10），不是单位名。

    直接把它当单位会把 unit 存成 '50'、'9.999999999999998' 这种数字，
    库存显示就变成"50 50"这种鬼东西。这里判断：整列都是数字 -> 不是单位列，
    真实单位从表头括号里取，「单位（卷）」的"卷"才是单位。
    """
    import re as _re
    has_num = False
    for v in sample_vals:
        if v is None or str(v).strip() == '':
            continue
        try:
            float(str(v)); has_num = True; break
        except (TypeError, ValueError):
            return None                      # 有文本 -> 确实就是单位列，原样返回
    if not has_num:
        return None
    m = _re.search(r'[（(]([^）)]+)[)）]', str(label or ''))
    if m:
        u = m.group(1).strip()
        # 括号里可能是"卷""平米""KG"等单位名，不该是"米""㎡"这种量纲
        if u and len(u) <= 6 and not _re.match(r'^[\d.]+$', u):
            return u
    return None

def map_headers(header, aliases=None):
    """表头行 -> {列索引: 规范字段}。aliases 来自列映射配置，优先级高于内置别名"""
    names = {f: list(v) for f, v in ALIASES.items()}
    for f, v in (aliases or {}).items():
        if f in names:
            names[f] = list(v) + [x for x in ALIASES.get(f, []) if x not in v]
    out = {}
    for i, h in enumerate(header):
        n = _norm(h)
        if not n:
            continue
        for f, ns in names.items():
            if f in out.values():
                continue
            if n in [_norm(x) for x in ns]:
                out[i] = f
                break
    return out

def _pick_sheet(cand, aliases=None):
    """从多张表候选里挑出真正有数据的那张。

    两轮：先找表头能被识别的（最可靠），
    都不认识就退回第一张非空表（至少让用户看到诊断网格去手选列）。
    """
    for rows in cand:
        for i in range(min(5, len(rows))):
            try:
                if map_headers(rows[i], aliases):
                    return rows
            except Exception:
                pass
    return cand[0]


def parse_file(path=None, stream=None, filename='', aliases=None):
    """返回 (行列表[dict], 识别到的字段集合)"""
    if path and path.lower().endswith('.csv') or (filename or '').lower().endswith('.csv'):
        data = open(path, 'rb').read() if path else stream.read()
        for enc in ('utf-8-sig', 'gbk', 'utf-8'):
            try:
                text = data.decode(enc); break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode('utf-8', 'replace')
        rows = list(csv.reader(io.StringIO(text)))
    else:
        if openpyxl is None:
            raise RuntimeError('未安装 openpyxl，请先执行: pip install openpyxl')
        wb = openpyxl.load_workbook(path or io.BytesIO(stream.read()), data_only=True)
        # 选表：很多人第一张表是封面/说明，真正的数据在后面。
        # 直接取 sheetnames[0] 会停在封面上，表头认不出、导入 0 条，
        # 用户只会看到"没认出表头"，完全想不到是选错了表。
        # 策略：先找表头能认出来的那张；都不认识才退回第一张非空表。
        _cand = []
        for _nm in wb.sheetnames:
            _rows = [[c for c in r] for r in wb[_nm].iter_rows(values_only=True)]
            _rows = [r for r in _rows if any(_norm(c) for c in r)]
            if _rows:
                _cand.append(_rows)
        rows = _pick_sheet(_cand, aliases) if _cand else []

    rows = [r for r in rows if any(_norm(c) for c in r)]
    if not rows:
        return [], set()

    hi, hmap = 0, {}
    for i in range(min(5, len(rows))):
        m = map_headers(rows[i], aliases)
        if 'name' in m.values() and len(m) > len(hmap):
            hi, hmap = i, m
    if 'name' not in hmap.values():          # 无表头：A=名称 B=料号 C=类型 D=单位 E=期初
        hmap = {0: 'name', 1: 'code', 2: 'category', 3: 'unit', 4: 'opening'}
        # 首行若含任何已知字段名，仍当作表头跳过，避免把表头当数据
        known = set()
        for v in (aliases or {}).values():
            known |= {_norm(x) for x in v}
        for v in ALIASES.values():
            known |= {_norm(x) for x in v}
        hi = 0 if any(_norm(c) in known for c in rows[0]) else -1

    out = []
    for r in rows[hi + 1:]:
        d = {f: '' for f in ALIASES}
        for i, f in hmap.items():
            if i < len(r):
                v = r[i]
                d[f] = '' if v is None else _as_text(v)
        d['name'] = d['name'] or d['code']
        if not d['name']:
            continue
        for f in NUM:
            d[f] = _num(d[f], hi=QTY_MAX)
        out.append(d)
    return out, set(hmap.values())

def _as_text(v):
    """单元格值转文本。

    坑：宽幅常见值是 0.035/0.045/0.05 这类真小数，以前用
    str(v).split('.')[0] 想去掉 250.0 的尾巴，结果把 0.035 也砍成了 "0"
    —— 宽幅信息全丢，连带"名称+规格+宽幅"的物料指纹失效，
    不同宽幅的物料被合并成一条，库存张冠李戴。
    正确做法：只有整数值（250.0）才去掉小数部分，真小数原样保留。
    """
    if v is None:
        return ''
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:      # 250.0 -> "250"
            return str(int(v))
        # 去掉浮点噪音：0.035000000000000003 -> 0.035
        return ('%g' % v)
    return str(v).strip()


def _num(v, lo=None, hi=None):
    """转数字。lo/hi 为边界，超出则截断；非法值一律 0。

    数量上限很必要：表里填错多几个零（1e15）会把库存冲到天文数字，
    后面所有报表、预警全都失真，而且很难一眼看出是哪笔错了。
    """
    try:
        f = float(str(v).strip().replace(',', ''))
    except (TypeError, ValueError):
        return 0.0
    import math
    if math.isnan(f) or math.isinf(f):
        return 0.0
    if lo is not None and f < lo:
        f = lo
    if hi is not None and f > hi:
        f = hi
    return f

def _read_grid(path=None, stream=None, filename=''):
    """把 xlsx/xlsm/xls/csv 读成二维列表。

    选表：很多人第一张表是封面/说明，真正的数据在后面。
    以前直接取"第一张非空表"，但封面通常也有几个字，照样会选中它，
    结果表头认不出、导入 0 条，用户只会看到"没认出表头"，
    完全想不到是程序选错了表。
    现在分两轮：先找表头能被认出来的那张；都不认识才退回第一张非空表
    （至少让用户看到诊断网格，可以手动指定列）。
    """
    fn = ((path or '') + (filename or '')).lower()
    if fn.endswith('.csv'):
        data = open(path, 'rb').read() if path else stream.read()
        for enc in ('utf-8-sig', 'gbk', 'utf-8', 'gb18030'):
            try:
                text = data.decode(enc); break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode('utf-8', 'replace')
        return list(csv.reader(io.StringIO(text)))
    if fn.endswith(('.xls', '.et')):
        try:
            import xlrd
        except ImportError:
            raise RuntimeError('读取 .xls/.et 需要 xlrd，请先执行 pip install xlrd，'
                               '或在 WPS 里另存为 .xlsx / .csv 再导入')
        bk = xlrd.open_workbook(path or (io.BytesIO(stream.read()) if stream else None))
        sh = bk.sheet_by_index(0)
        return [[_cell(sh, r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    if openpyxl is None:
        raise RuntimeError('未安装 openpyxl，请先执行: pip install openpyxl')
    wb = openpyxl.load_workbook(path or io.BytesIO(stream.read()), data_only=True)
    cand = []
    for nm in wb.sheetnames:
        rows = [[c for c in r] for r in wb[nm].iter_rows(values_only=True)]
        rows = [r for r in rows if any(_norm(c) for c in r)]
        if rows:
            cand.append(rows)
    if not cand:
        return []
    for rows in cand:                       # 第一轮：表头可识别的优先
        for i in range(min(5, len(rows))):
            try:
                if map_txn_headers(rows[i]):
                    return rows
            except Exception:
                pass
    return cand[0]                          # 第二轮：退回第一张非空表


def _cell(sh, r, c):
    v = sh.cell_value(r, c)
    if sh.cell_type(r, c) == 3:            # XL_CELL_DATE
        import xlrd, datetime as _dt
        try:
            t = xlrd.xldate_as_tuple(v, sh.book.datemode)
            return _dt.datetime(*t).strftime('%Y-%m-%d')
        except Exception:
            return str(v)
    return v

def _fmt_date(v):
    """把 Excel 日期/字符串统一成 YYYY-MM-DD"""
    if v is None or str(v).strip() == '':
        return ''
    if isinstance(v, (int, float)):        # Excel 序列号
        from datetime import date as _d, timedelta
        try:
            return (_d(1899, 12, 30) + timedelta(days=float(v))).strftime('%Y-%m-%d')
        except Exception:
            return str(v)
    s = str(v).strip()
    for attr in ('date', 'datetime'):
        pass
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    # 只做"整体像日期"的替换：先确认格式，再动手，避免把文本里的"日"字吃掉
    import re as _re
    m = _re.match(r'^\s*(\d{4})\s*[-/年.]\s*(\d{1,2})\s*[-/月.]\s*(\d{1,2})\s*日?\s*$', s)
    if m:
        y, mo, d = m.group(1), m.group(2).zfill(2), m.group(3).zfill(2)
        try:
            datetime.date(int(y), int(mo), int(d))   # 顺带校验合法性（排除 2026-13-45）
        except ValueError:
            return ''
        return f'{y}-{mo}-{d}'
    m2 = _re.match(r'^\s*(\d{1,2})\s*[-/月.]\s*(\d{1,2})\s*日?\s*$', s)   # 只有月日
    if m2:
        return ''
    # 认不出来就返回空，由上层用默认日期，绝不把垃圾塞进数据库
    return ''

def map_txn_headers(header, aliases=None, mat_aliases=None):
    """流水表头 -> {列索引: 字段}。先看进出/数量专用别名，再补物料 A-G 列"""
    names = {f: list(v) for f, v in TXN_ALIASES.items()}
    for f, v in (aliases or {}).items():          # 用户在列映射里自定义的
        if f in names:
            names[f] = list(v) + [x for x in TXN_ALIASES.get(f, []) if x not in v]
    mat = dict(ALIASES)
    for f, v in (mat_aliases or {}).items():
        if f in mat:
            mat[f] = list(v) + [x for x in ALIASES.get(f, []) if x not in v]
    out = {}
    # 先匹配流水专用列（避免「数量」被物料列抢走）
    for i, h in enumerate(header):
        n = _norm(h)
        if not n:
            continue
        for f in ('date', 'kind', 'qty', 'in_qty', 'out_qty', 'pieces', 'per_piece',
                  'price', 'amount', 'note'):
            if f in out.values():
                continue
            if n in [_norm(x) for x in names[f]]:
                out[i] = f; break
    # 再匹配 name/code 与物料 A-G 列
    for i, h in enumerate(header):
        if i in out:
            continue
        n = _norm(h)
        if not n:
            continue
        for f in ['name', 'code'] + TXN_MAT_COLS:
            if f in out.values():
                continue
            # names（流水内置别名）和 mat（用户自定义别名）要合起来看。
            # 原来写成 names.get(f) or mat.get(f)：name/code 在 names 里本来就有，
            # 于是用户自己起的表头名永远轮不到，改名后导入就认不出来了。
            pool = list(names.get(f) or [])
            for x in (mat.get(f) or []):
                if x not in pool:
                    pool.append(x)
            if n in [_norm(x) for x in pool]:
                out[i] = f; break
    return out

# ---------- 原 Excel「宽表」布局：物料一行 × 每日(进/出)两列 ----------
WIDE_KINDS = {'进', '出'}
MAT_HEADS = {
    'supplier': ['供应商', '厂商', '供货商'],
    'category': ['类型', '类别', '分类', '大类'],
    'spec': ['规格', '规格(米)', '规格（米）', '厚度'],
    'width': ['宽幅', '宽度', '幅宽'],
    'name': ['物料名称', '名称', '品名', '材料名称', '物料'],
    'code': ['料号', '物料编号', '物料编码', '编号', '编码', '型号'],
    'unit': ['单位', '单位(卷)', '单位（卷）', '计量单位'],
    'opening': ['上月结存', '期初结存', '期初', '上期结存'],
}

def _is_dt(v):
    return hasattr(v, 'strftime') or isinstance(v, datetime.datetime)

def _wide_kind_row(grid):
    """找形如 [进,出,进,出...] 的那一行（原表第2行）"""
    best, bestn = None, 0
    for r in range(min(8, len(grid))):
        n = sum(1 for v in grid[r] if str(v or '').strip() in WIDE_KINDS)
        if n > bestn:
            best, bestn = r, n
    return best if bestn >= 4 else None

def _head_at(grid, rows, c):
    """取第 c 列在若干表头行里的文字（优先非空的那行）"""
    for hr in rows:
        if hr is not None and hr < len(grid) and c < len(grid[hr]):
            v = str(grid[hr][c] or '').strip()
            if v:
                return v
    return ''

# 「汇总/合计」这类列长在日期区后面，表头不是日期、进出标记却同样是「进/出」。
# 当作日期列读进来会把当月的合计值再导一遍 —— 库存直接翻倍。必须整段跳过。
SUM_WORDS = ('汇总', '合计', '总计', '小计', '累计', '合計')

def _is_sum_col(v):
    s = str(v or '').strip()
    return bool(s) and any(w in s for w in SUM_WORDS)

def _wide_dates(grid, drow, defmonth):
    """从日期行取每列日期；合并单元格造成的 None 沿用前一个日期。

    遇到「汇总/合计」列立刻中断日期延续（last=None），
    否则它右边的列会继续沿用汇总列之前那个日期，同样被误当成单据列。
    """
    row = grid[drow] if (drow is not None and drow < len(grid)) else []
    y, mo = (int(defmonth[:4]), int(defmonth[5:7])) if defmonth and len(defmonth) >= 7 \
        else (datetime.date.today().year, datetime.date.today().month)
    out, last = {}, None
    for c, v in enumerate(row):
        if _is_sum_col(v):
            last = None
            continue
        d = None
        if _is_dt(v):
            d = v.strftime('%Y-%m-%d')
        elif isinstance(v, (int, float)) and 1 <= float(v) <= 31:
            d = '%04d-%02d-%02d' % (y, mo, int(v))
        else:
            s = str(v or '').strip()
            if s and (('月' in s) or ('-' in s) or ('/' in s)):
                d = _fmt_date(v) or None
        if d:
            last = d
        if last:
            out[c] = last
    return out

def _strip_paren(n):
    """去掉表头末尾的括号单位：「宽幅(米)」->「宽幅」、「单位(卷)」->「单位」。

    真实表格里大量写成「规格（米）」「宽幅（米）」「单位（卷）」，
    只按全名精确匹配会整列丢失（宽幅直接变空）。
    """
    import re as _re
    return _re.sub(r'[（(][^）)]*[)）]\s*$', '', n).strip()

def _col_by_head(grid, hrows, names):
    """在若干候选表头行里按名字找列号。

    先精确匹配，再退一步去掉括号后缀匹配（「宽幅(米)」按「宽幅」命中）。
    """
    exact = [_norm(x) for x in names]
    loose = [_strip_paren(x) for x in exact]
    for hr in hrows:
        if hr is None or hr >= len(grid):
            continue
        # 先跑一遍精确，保证「规格」不会被「规格(米)」抢走
        for c, v in enumerate(grid[hr]):
            n = _norm(v)
            if n and n in exact:
                return c
    for hr in hrows:
        if hr is None or hr >= len(grid):
            continue
        for c, v in enumerate(grid[hr]):
            n = _strip_paren(_norm(v))
            if n and n in loose:
                return c
    return None

def _combo_date_kind(v, y, mo):
    """识别「日期+进出」写在同一个单元格里的表头。

    真实表格里这种写法很常见：09-01进 / 09-01出 / 9/1 进 / 2026-09-01进。
    以前只认"一行日期 + 一行进出"的双层表头，遇到这种单层写法
    parse_wide 直接返回空，用户又没法手动指定列（宽表列太多，
    手动映射只适合一行一笔的表），数据就彻底导不进来了。
    """
    s = str(v or '').strip()
    if not s:
        return None
    kind = None
    for k in WIDE_KINDS:
        if s.endswith(k):
            kind = k
            s = s[:-1].strip()
            break
    if kind is None:
        return None
    s = s.rstrip('/-—– 　')
    if not s:
        return None
    d = _fmt_date(s)                       # 完整日期：2026-09-01 / 2026/9/1
    if d:
        return (d, kind)
    m = re.match(r'^(\d{1,2})\s*[\-/月]\s*(\d{1,2})\s*日?$', s)   # 09-01 / 9-1 / 9月1
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        mm, dd = (a, b) if a <= 12 else (b, a)
        if 1 <= mm <= 12 and 1 <= dd <= 31:
            try:
                return ('%04d-%02d-%02d' % (y, mo, dd) if mm == mo
                        else '%04d-%02d-%02d' % (y, mm, dd), kind)
            except ValueError:
                return None
    return None


def _combo_cols(grid, defmonth=None):
    """扫描表头区，找出"日期+进出"合并写法的列。

    返回 (dates, kind_map, 最后一条表头行号)；找不到返回 None。
    """
    y, mo = (int(defmonth[:4]), int(defmonth[5:7])) if defmonth and len(defmonth) >= 7 \
        else (datetime.date.today().year, datetime.date.today().month)
    dates, kinds, last = {}, {}, None
    for r in range(min(6, len(grid))):
        for c, v in enumerate(grid[r] if r < len(grid) else []):
            if c in dates:
                continue
            got = _combo_date_kind(v, y, mo)
            if got:
                dates[c], kinds[c] = got
                last = r if last is None else max(last, r)
    if len(dates) >= 4 and kinds:          # 至少 4 个日期列才认定是宽表
        return dates, kinds, last
    return None


def parse_wide(grid, defmonth=None):
    """解析原表那种『物料 × 每日进/出』宽表。返回 (单据列表, 字段集合)

    支持两种表头写法：
      A 双层：一行写日期、下一行写 进/出（原表就是这种）
      B 单层：日期和进出写在一个格子里，如「09-01进」「09-01出」
    """
    kr = _wide_kind_row(grid)
    combo = None
    if kr is None:
        combo = _combo_cols(grid, defmonth)     # 试试单层合并写法
        if combo is None:
            return [], set()
        dates, kind_map_pre, last = combo
        start = (last or 0) + 1
        heads = [r for r in range((last or 0) + 1)]
        kr = None
    else:
        drow = kr - 1 if kr > 0 else None
        dates = _wide_dates(grid, drow, defmonth)
        if not dates:
            return [], set()
        # 数据起点：进出行之后，且该行有物料名称
        start = kr + 1
        heads = [hr for hr in (drow, kr) if hr is not None]
        kind_map_pre = None
    name_c = _col_by_head(grid, heads, MAT_HEADS['name'])
    if name_c is None:                      # 表头对不上：挑"文字最多"的那一列当物料名
        best, bestn = None, 0
        for c in range(min(12, max(len(r) for r in grid[:20]))):
            n = 0
            for r in grid[start:start + 15]:
                if r and c < len(r):
                    v = str(r[c] or '').strip()
                    if v and not _num(v) and len(v) >= 2:
                        n += 1
            if n > bestn:
                best, bestn = c, n
        name_c = best
    if name_c is None:
        return [], set()
    cols = {k: _col_by_head(grid, heads, v) for k, v in MAT_HEADS.items()}
    cols['name'] = name_c
    # 日期列 -> (列号, 进/出)
    if kind_map_pre is not None:
        kind_map = dict(kind_map_pre)        # 单层写法：方向直接来自表头单元格
    else:
        kind_map = {}
        for c, v in enumerate(grid[kr] if (kr is not None and kr < len(grid)) else []):
            s = str(v or '').strip()
            if s in WIDE_KINDS and c in dates:
                kind_map[c] = s
    out = []
    for r in grid[start:]:
        if not r or name_c >= len(r):
            continue
        nm = str(r[name_c] or '').strip()
        if not nm or nm in ('/', '-', '合计', '总计'):
            continue
        base = {'name': nm, 'code': '', 'date': '', 'note': '',
                'pieces': None, 'per_piece': None, 'price': None}
        for f in ('supplier', 'category', 'spec', 'width', 'unit', 'code'):
            c = cols.get(f)
            base[f] = '' if (c is None or c >= len(r)) else str(r[c] or '').strip()
            if base[f] in ('/', '-'):
                base[f] = ''
        # 「单位（卷）」列里填的是卷数，不是单位名；把真实单位从表头括号里取出来
        uc = cols.get('unit')
        if uc is not None:
            fixed = unit_from_label(_head_at(grid, heads, uc),
                                    [r[uc] if uc < len(r) else None for r in grid[start:start + 30]])
            if fixed:
                base['unit'] = fixed
        oc = cols.get('opening')
        base['opening'] = _num(r[oc], hi=QTY_MAX) if (oc is not None and oc < len(r)) else 0.0
        base['safety'] = 0.0
        got = False
        for c, kind in sorted(kind_map.items()):
            if c >= len(r):
                continue
            q = _num(r[c], hi=QTY_MAX)
            if q > 0:
                d = dict(base); d['kind'] = kind; d['qty'] = q
                d['date'] = dates[c]
                out.append(d); got = True
        if not got:                          # 该行没有进出，但档案信息仍要保留（期初）
            d = dict(base); d['kind'] = None; d['qty'] = 0.0
            out.append(d)
        # 数量取整到合理精度：原表里 9.999999999999998 这种浮点尾巴
        # 会一路带进数据库，显示成"9.999999999999998 卷"
    for d in out:
        if d.get('qty'):
            d['qty'] = round(float(d['qty']), 6)
    return out, set(list(cols.keys()) + ['date', 'qty', 'kind'])

def parse_txn_by_map(path=None, stream=None, filename='', colmap=None, start=1,
                     default_kind=None, defdate=None):
    """按用户手动指定的列来解析。colmap: {列索引: 字段名}，start: 数据起始行(0基)
    字段名取值 name/code/date/kind/qty/in_qty/out_qty/pieces/per_piece/
    price/amount/note，未列出的列忽略。
    金额(amount)可以不填，由 数量×单价 自动算出；只填金额时反推单价。"""
    grid = _read_grid(path=path, stream=stream, filename=filename)
    colmap = {int(k): v for k, v in (colmap or {}).items() if v}
    out = []
    for r in grid[start:]:
        if not any(_norm(c) for c in r):
            continue
        def val(f):
            for k, v in colmap.items():
                if v == f:
                    return r[k] if k < len(r) else None
            return None
        name = str(val('name') or '').strip()
        code = str(val('code') or '').strip()
        if not name and not code:
            continue
        base = dict(name=name or code, code=code,
                    date=_fmt_date(val('date')) or (defdate or ''),
                    note=str(val('note') or '').strip(),
                    pieces=_num(val('pieces'), hi=QTY_MAX) or None,
                    per_piece=_num(val('per_piece'), hi=QTY_MAX) or None)
        for f in TXN_MAT_COLS:
            v = val(f)
            base[f] = '' if v is None else str(v).strip()
            if f in ('opening', 'safety'):
                base[f] = _num(v)
        ins = _num(val('in_qty'), hi=QTY_MAX); outs = _num(val('out_qty'), hi=QTY_MAX)
        q = _num(val('qty'), hi=QTY_MAX)
        pcs = _num(val('pieces'), hi=QTY_MAX); per = _num(val('per_piece'), hi=QTY_MAX)
        if not q and pcs and per:
            q = round(pcs * per, 2)
        # 单价：表里直接给"单价"最好；只给"金额"时按 金额÷数量 反推。
        # 两个都给了以单价为准（金额只是展示值，不入库，避免对不上）。
        pr = _num(val('price'))
        if not pr:
            am = _num(val('amount'))
            if am and q:
                pr = round(am / q, 4)
        pr = pr or None
        kd = str(val('kind') or '').strip()
        items = []
        if ins or outs:
            if ins: items.append(('进', ins))
            if outs: items.append(('出', outs))
        elif q:
            k = default_kind or ''
            if kd:
                k = '出' if ('出' in kd or 'out' in kd.lower()) else '进'
            items.append((k or '进', q))
        elif per and pcs:
            pass
        for kind, qty in items:
            if qty <= 0:
                continue
            d = dict(base); d['kind'] = kind; d['qty'] = qty; d['price'] = pr
            out.append(d)
    return out

def sniff_grid(grid, n=6, w=14):
    """给诊断用：原始前几行前几列"""
    return [[('' if c is None else str(c))[:18] for c in (r or [])[:w]]
            for r in grid[:n]]

def parse_txn_file(path=None, stream=None, filename='', aliases=None, mat_aliases=None,
                   default_kind=None, defmonth=None, xcols=None):
    """解析流水表。先按"一行一笔"解析；失败则按原表宽表解析。
    返回 (单据列表, 识别到的字段集合, 表头行号, 诊断信息)"""
    grid = _read_grid(path=path, stream=stream, filename=filename)
    raw = [r for r in grid if any(_norm(c) for c in r)]
    if not raw:
        return [], set(), -1, {'why': '空文件', 'grid': []}

    hi, hmap = 0, {}
    for i in range(min(5, len(raw))):
        m = map_txn_headers(raw[i], aliases, mat_aliases)
        has_q = ('qty' in m.values() or 'in_qty' in m.values() or 'out_qty' in m.values()
                 or ('pieces' in m.values() and 'per_piece' in m.values()))
        if 'name' in m.values() and has_q and len(m) > len(hmap):
            hi, hmap = i, m

    if 'name' not in hmap.values():
        # 不是"一行一笔"，试试原表那种"物料 × 每日进/出"宽表
        wrows, wfields = parse_wide(raw, defmonth)
        skipped = ([c for c, v in enumerate(raw[0]) if _is_sum_col(v)]
                   if raw else [])
        if wrows:
            # 原表常把下月 1 号也画进来（"9月表"里含 10-01 列），
            # 不提示的话用户会以为只导了本月，回头对不上账
            yms = {}
            for r in wrows:
                ym = str(r.get('date') or '')[:7]
                if ym: yms[ym] = yms.get(ym, 0) + 1
            return wrows, wfields, -2, {'why': 'wide', 'grid': sniff_grid(raw),
                                        'yms': yms, 'skipped_sum': len(skipped)}
        return [], set(), -1, {
            'why': 'no-header',
            'grid': sniff_grid(raw),
            'heads': [str(v or '') for v in (raw[0] if raw else [])][:14],
        }

    # 自定义列（客户订单号之类）：表头里能按列名对上的，一并读进来
    if xcols:
        for i, h in enumerate(raw[hi]):
            if i in hmap:
                continue
            n = _norm(h)
            if not n:
                continue
            for x in xcols:
                if n == _norm(x['label']):
                    hmap[i] = x['fid']
                    break

    out = []
    for r in raw[hi + 1:]:
        def val(f):
            for k, v in hmap.items():
                if v == f:
                    return r[k] if k < len(r) else None
            return None
        name = str(val('name') or '').strip()
        code = str(val('code') or '').strip()
        if not name and not code:
            continue
        base = dict(name=name or code, code=code,
                    date=_fmt_date(val('date')),
                    note=str(val('note') or '').strip(),
                    pieces=_num(val('pieces'), hi=QTY_MAX) or None,
                    per_piece=_num(val('per_piece'), hi=QTY_MAX) or None)
        for f in TXN_MAT_COLS:
            v = val(f)
            base[f] = '' if v is None else str(v).strip()
            if f in ('opening', 'safety'):
                base[f] = _num(v, hi=QTY_MAX)
        # 自定义列的值：hmap 里已按列名对上，这里原样带出去
        for k, f in hmap.items():
            if str(f).startswith('x_') and k < len(r):
                base[f] = '' if r[k] is None else str(r[k]).strip()
        # 数量：优先 进/出 分列，其次 数量 + 进出方向
        # 都加 QTY_MAX 上限：表里多输几个零不该把库存冲到天文数字
        ins = _num(val('in_qty'), hi=QTY_MAX)
        outs = _num(val('out_qty'), hi=QTY_MAX)
        q = _num(val('qty'), hi=QTY_MAX)
        pcs = _num(val('pieces'), hi=QTY_MAX)
        per = _num(val('per_piece'))
        if not q and pcs and per:      # 只有件数×每件 -> 总数量
            q = round(pcs * per, 2)
        # 单价：优先表里的"单价"列；只给"金额"时按 金额÷数量 反推
        pr = _num(val('price'))
        if not pr:
            am = _num(val('amount'))
            if am and q:
                pr = round(am / q, 4)
        pr = pr or None
        base['qty_computed'] = bool(pcs and per and not _num(val('qty')))
        kd = str(val('kind') or '').strip()
        items = []
        if ins or outs:
            if ins:
                items.append(('进', ins))
            if outs:
                items.append(('出', outs))
        elif q:
            k = default_kind or ''
            if kd:
                k = '出' if ('出' in kd or 'out' in kd.lower() or kd in ('-', 'OUT')) else '进'
            items.append((k or '进', q))
        for kind, qty in items:
            if qty <= 0:
                continue
            d = dict(base); d['kind'] = kind; d['qty'] = qty; d['price'] = pr
            out.append(d)
    if not out:
        return out, set(hmap.values()), hi, {
            'why': 'no-data', 'grid': sniff_grid(raw),
            'heads': [str(v or '') for v in raw[hi]][:14],
            'hi': hi,
        }
    return out, set(hmap.values()), hi, {'why': 'long', 'hi': hi}
